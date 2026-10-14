"""Coleta e renderização de relatórios — ETAPA 16 (CSV, XLSX, PDF).

Coleta reutiliza os repositórios dos módulos (mesmos filtros das telas).
Renderizadores genéricos: (colunas, linhas) → bytes.
"""

import csv
from io import BytesIO, StringIO

FORMATOS = {"csv", "xlsx", "pdf"}
RECURSOS = {"funcionarios", "presencas", "apropriacoes", "horas-extras", "pendencias", "tarefas", "absenteismo", "emails",
            "preventivas-programacao", "preventivas-execucao", "preventivas-pendencias"}
LIMITE_PDF = 500


def _nomes(func_ids: set) -> dict:
    if not func_ids:
        return {}
    from app.core.supabase_client import get_supabase

    sb = get_supabase()
    r = sb.table("funcionarios").select("id,nome,matricula").in_("id", sorted(func_ids)).execute()
    return {x["id"]: f"{x.get('matricula', '')} — {x.get('nome', '')}" for x in (r.data or [])}


def coletar(recurso: str, f: dict) -> tuple[list, list]:
    if recurso == "funcionarios":
        from app.modules.funcionarios import repository as r

        rows = r.list_repo(f.get("q"), f.get("ativo"), f.get("setor"), None, None, 1, 500)["items"]
        cols = ["matricula", "nome", "cargo", "setor", "empresa", "admissao", "ativo"]
        return cols, [[x.get(c) for c in cols] for x in rows]
    if recurso == "presencas":
        from app.modules.presencas import repository as r

        rows = r.list_repo(f.get("data"), f.get("de"), f.get("ate"), f.get("funcionario_id"), f.get("status"))
        cols = ["data", "funcionario", "status_codigo", "obs"]
        out = []
        for x in rows:
            fn = x.get("funcionarios") or {}
            out.append([x.get("data"), fn.get("nome") or x.get("funcionario_id"), x.get("status_codigo"), x.get("obs")])
        return cols, out
    if recurso == "apropriacoes":
        from app.modules.apropriacoes import repository as r

        rows = r.list_repo(f.get("data"), f.get("de"), f.get("ate"), f.get("funcionario_id"), None)
        nomes = _nomes({x["funcionario_id"] for x in rows})
        cols = ["data", "funcionario", "os", "horas"]
        return cols, [[x.get("data"), nomes.get(x["funcionario_id"], x["funcionario_id"]),
                       (x.get("ordens_servico") or {}).get("codigo"), x.get("horas")] for x in rows]
    if recurso == "horas-extras":
        from app.modules.horas_extras import repository as r

        rows = r.list_repo(f.get("data"), f.get("de"), f.get("ate"), f.get("funcionario_id"), None)
        nomes = _nomes({x["funcionario_id"] for x in rows})
        cols = ["data", "funcionario", "qtd_horas", "percentual"]
        return cols, [[x.get("data"), nomes.get(x["funcionario_id"], x["funcionario_id"]),
                       x.get("qtd_horas"), x.get("percentual")] for x in rows]
    if recurso == "pendencias":
        from app.modules.pendencias import repository as r

        rows = r.list_repo(f.get("status"), f.get("tipo"), f.get("funcionario_id"), f.get("de"), f.get("ate"))
        nomes = _nomes({x["funcionario_id"] for x in rows if x.get("funcionario_id")})
        cols = ["tipo", "data_ref", "funcionario", "descricao", "prioridade", "status"]
        return cols, [[x.get("tipo"), x.get("data_ref"), nomes.get(x.get("funcionario_id"), x.get("funcionario_id")),
                       x.get("descricao"), x.get("prioridade"), x.get("status")] for x in rows]
    if recurso == "tarefas":
        from app.modules.tarefas import repository as r

        rows = r.list_repo(f.get("status"), f.get("responsavel_id"), None, f.get("q"))
        nomes = _nomes({x["responsavel_id"] for x in rows if x.get("responsavel_id")})
        cols = ["titulo", "status", "prioridade", "responsavel", "prazo", "categoria"]
        return cols, [[x.get("titulo"), x.get("status"), x.get("prioridade"),
                       nomes.get(x.get("responsavel_id"), x.get("responsavel_id")),
                       x.get("prazo"), x.get("categoria")] for x in rows]
    if recurso == "absenteismo":
        from app.modules.absenteismo import repository as r

        rows = r.list_repo(f.get("de"), f.get("ate"), f.get("funcionario_id"), None, f.get("setor"))
        cols = ["data", "funcionario", "tipo", "periodo", "motivo"]
        out = []
        for x in rows:
            fn = x.get("funcionarios") or {}
            tp = x.get("tipos_ausencia") or {}
            out.append([x.get("data"), fn.get("nome") or x.get("funcionario_id"),
                        tp.get("codigo") or x.get("tipo_id"), x.get("periodo"), x.get("motivo")])
        return cols, out
    if recurso == "emails":
        from app.modules.emails import repository as r

        rows = r.list_repo(f.get("status"), f.get("de"), f.get("ate"))
        cols = ["assunto", "destinatario", "status", "enviado_em", "erro"]
        return cols, [[x.get(c) for c in cols] for x in rows]
    if recurso.startswith("preventivas-"):
        return _coletar_preventivas(recurso, f)
    raise KeyError(f"Recurso inválido: {recurso}")


def _coletar_preventivas(recurso: str, f: dict) -> tuple[list, list]:
    from app.modules.preventivas import service as pv

    itens = pv.listar(
        data=f.get("data"), de=f.get("de"), ate=f.get("ate"), area=f.get("area"),
        status=f.get("status"), funcionario_id=f.get("funcionario_id"),
        cargo=f.get("cargo"), os_codigo=f.get("os_codigo"), situacao=f.get("situacao"))
    detalhes = [pv.montar_detalhe(i["id"]) for i in itens]
    usuarios = _usuarios_nomes()
    if recurso == "preventivas-programacao":
        cols = ["Data", "Area", "Equipamento", "Preventiva", "Funcionario", "Matricula",
                "Cargo", "Escala", "Prev. inicio", "Prev. fim", "Responsavel", "Situacao"]
        rows = []
        for d in detalhes:
            p = d["preventiva"]
            for l in d["funcionarios"]:
                fn = l["funcionario"]
                rows.append([p.get("data"), p.get("area"), p.get("equipamento"), p.get("descricao"),
                             fn.get("nome"), fn.get("matricula"), fn.get("cargo"), fn.get("escala"),
                             l.get("hora_prev_inicio"), l.get("hora_prev_fim"),
                             p.get("responsavel_nome") or "", l.get("status")])
        return cols, rows
    if recurso == "preventivas-execucao":
        cols = ["Data", "Area", "Preventiva", "Funcionario", "Matricula", "Compareceu",
                "Efet. inicio", "Efet. fim", "OS", "Descricao OS", "Horas", "Recebida",
                "Data receb.", "Recebido por"]
        rows = []
        for d in detalhes:
            p = d["preventiva"]
            for l in d["funcionarios"]:
                if not l["compareceu"]:
                    continue
                fn = l["funcionario"]
                if l["os"]:
                    for o in l["os"]:
                        rows.append([p.get("data"), p.get("area"), p.get("descricao"),
                                     fn.get("nome"), fn.get("matricula"), "Sim",
                                     l.get("hora_efet_inicio"), l.get("hora_efet_fim"),
                                     o.get("os_codigo"), o.get("descricao"), o.get("horas"),
                                     "Sim" if o.get("recebida") else "Nao", o.get("data_recebimento"),
                                     usuarios.get(str(o.get("recebido_por")), "")])
                else:
                    rows.append([p.get("data"), p.get("area"), p.get("descricao"),
                                 fn.get("nome"), fn.get("matricula"), "Sim",
                                 l.get("hora_efet_inicio"), l.get("hora_efet_fim"),
                                 "-", "-", 0, "-", "-", ""])
        return cols, rows
    # preventivas-pendencias
    cols = ["Data", "Area", "Funcionario", "Matricula", "OS pendente",
            "OS recebidas", "OS pendentes", "Situacao"]
    rows = []
    for d in detalhes:
        p = d["preventiva"]
        for l in d["funcionarios"]:
            if l["status"] not in ("AGUARDANDO_OS", "OS_PENDENTE", "PARCIALMENTE_REGULARIZADO", "NAO_COMPARECEU"):
                continue
            fn = l["funcionario"]
            pendentes = [o.get("os_codigo") for o in l["os"] if not o.get("recebida")]
            rows.append([p.get("data"), p.get("area"), fn.get("nome"), fn.get("matricula"),
                         ", ".join([str(x) for x in pendentes]) or "-",
                         l["os_recebidas"], l["os_pendentes"], l["status"]])
    return cols, rows


def _usuarios_nomes() -> dict:
    from app.core.supabase_client import get_supabase

    try:
        sb = get_supabase()
        r = sb.table("usuarios").select("id,nome").execute()
        return {str(x["id"]): x.get("nome") for x in (r.data or [])}
    except Exception:
        return {}


def render_csv(cols: list, rows: list) -> tuple[bytes, str]:
    buf = StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    w.writerows([["" if v is None else v for v in r] for r in rows])
    return buf.getvalue().encode("utf-8-sig"), "text/csv"


def render_xlsx(cols: list, rows: list) -> tuple[bytes, str]:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _txt(v) -> str:
    s = str(v if v is not None else "")
    return s.replace("—", "-").replace("–", "-").encode("latin-1", "replace").decode("latin-1")


def render_pdf(recurso: str, cols: list, rows: list) -> tuple[bytes, str]:
    from fpdf import FPDF

    rows = rows[:LIMITE_PDF]
    pdf = FPDF(orientation="L", format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, _txt(f"ERP - {recurso} ({len(rows)} linhas)"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 7)
    w = max(20, min(60, int(270 / max(len(cols), 1))))
    for c in cols:
        pdf.cell(w, 7, _txt(c)[:30], border=1)
    pdf.ln()
    for r in rows:
        for v in r:
            pdf.cell(w, 6, _txt(v)[:30], border=1)
        pdf.ln()
    return bytes(pdf.output()), "application/pdf"


def gerar(recurso: str, formato: str, filtros: dict) -> tuple[bytes, str, str]:
    if recurso not in RECURSOS:
        raise KeyError(f"Recurso inválido: {recurso}")
    if formato not in FORMATOS:
        raise ValueError(f"Formato inválido: {formato}")
    cols, rows = coletar(recurso, filtros)
    if formato == "csv":
        data, mime = render_csv(cols, rows)
    elif formato == "xlsx":
        data, mime = render_xlsx(cols, rows)
    else:
        data, mime = render_pdf(recurso, cols, rows)
    return data, mime, f"{recurso}.{formato}"
